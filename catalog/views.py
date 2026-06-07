import io
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth import login, authenticate, logout
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth import update_session_auth_hash
from django.core.paginator import Paginator
from django.db.models import Q
from django.contrib import messages
from django.core.mail import EmailMessage
from django.conf import settings
from openpyxl import Workbook
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from django.contrib.auth.models import User

from .models import Product, Category, Manufacturer, Cart, CartItem, Profile, Order, OrderItem
from .serializers import (
    CategorySerializer, ManufacturerSerializer, ProductSerializer, 
    CartSerializer, CartItemSerializer, UserSerializer, ProfileSerializer,
    OrderSerializer, ChangePasswordSerializer
)
from .permissions import IsAdminOrReadOnly, IsOwnerOrAdmin, IsAdminUser
from .forms import UserRegisterForm, ProfileUpdateForm

# ==================== HTML ВЕБ-ПРЕДСТАВЛЕНИЯ ====================

def home_view(request):
    """Главная страница магазина"""
    popular_products = Product.objects.all().order_by('-id')[:6]
    categories = Category.objects.all()
    return render(request, 'shop/index.html', {
        'popular_products': popular_products, 
        'categories': categories
    })

def product_list(request):
    """Страница каталога с фильтрацией и пагинацией"""
    products = Product.objects.all().order_by('id')
    categories = Category.objects.all()
    manufacturers = Manufacturer.objects.all()

    search_query = request.GET.get('search', '')
    category_id = request.GET.get('category', '')
    manufacturer_id = request.GET.get('manufacturer', '')

    if search_query:
        products = products.filter(
            Q(name__icontains=search_query) | 
            Q(description__icontains=search_query)
        )
    if category_id:
        products = products.filter(category_id=category_id)
    if manufacturer_id:
        products = products.filter(manufacturer_id=manufacturer_id)

    paginator = Paginator(products, 9)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'categories': categories,
        'manufacturers': manufacturers,
        'page_obj': page_obj,
        'search_query': search_query,
        'selected_category': category_id,
        'selected_manufacturer': manufacturer_id,
    }
    return render(request, 'shop/catalog.html', context)

def product_detail(request, pk):
    """Детальная страница товара"""
    product = get_object_or_404(Product, pk=pk)
    return render(request, 'shop/product_detail.html', {'product': product})

@login_required
def cart_view(request):
    """Просмотр корзины пользователя"""
    cart, created = Cart.objects.get_or_create(user=request.user)
    return render(request, 'shop/cart.html', {'cart': cart})

@login_required
def add_to_cart(request, product_id):
    """Добавление товара в корзину"""
    product = get_object_or_404(Product, id=product_id)
    cart, created = Cart.objects.get_or_create(user=request.user)
    cart_item, item_created = CartItem.objects.get_or_create(cart=cart, product=product)
    
    if not item_created:
        if cart_item.quantity + 1 <= product.stock_quantity:
            cart_item.quantity += 1
            cart_item.save()
            messages.success(request, f"Количество товара {product.name} увеличено.")
        else:
            messages.error(request, f"На складе осталось всего {product.stock_quantity} шт.")
    else:
        if product.stock_quantity >= 1:
            cart_item.quantity = 1
            cart_item.save()
            messages.success(request, f"Товар {product.name} добавлен в корзину.")
        else:
            messages.error(request, "Товара нет в наличии.")
    return redirect('product_list')

@login_required
def update_cart(request, item_id):
    """Обновление количества товара в корзине"""
    cart_item = get_object_or_404(CartItem, id=item_id, cart__user=request.user)
    if request.method == 'POST':
        try:
            new_quantity = int(request.POST.get('quantity', 1))
            if new_quantity <= 0:
                cart_item.delete()
            elif new_quantity <= cart_item.product.stock_quantity:
                cart_item.quantity = new_quantity
                cart_item.save()
                messages.success(request, "Количество обновлено.")
            else:
                messages.error(request, f"Доступно только {cart_item.product.stock_quantity} шт.")
        except ValueError:
            pass
    return redirect('cart_view')

@login_required
def remove_from_cart(request, item_id):
    """Удаление товара из корзины"""
    cart_item = get_object_or_404(CartItem, id=item_id, cart__user=request.user)
    cart_item.delete()
    messages.success(request, "Товар удален из корзины.")
    return redirect('cart_view')

@login_required
def checkout(request):
    """Оформление заказа"""
    cart, created = Cart.objects.get_or_create(user=request.user)
    cart_items = cart.items.all()
    if not cart_items:
        return redirect('cart_view')

    if request.method == 'POST':
        address = request.POST.get('address', '').strip()
        payment_method = request.POST.get('payment', 'card')
        delivery_method = request.POST.get('delivery', 'courier')
        
        if not address:
            messages.error(request, "Укажите адрес доставки.")
            return render(request, 'shop/checkout.html', {'cart': cart})

        order = Order.objects.create(
            user=request.user,
            address=address,
            total_price=cart.total_price,
            payment_method=payment_method,
            delivery_method=delivery_method,
            status='pending'
        )
        
        for item in cart_items:
            OrderItem.objects.create(
                order=order,
                product=item.product,
                quantity=item.quantity,
                price=item.product.price
            )
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Чек"
        ws['A1'] = f"Заказ #{order.id} пользователя: {request.user.username}"
        ws['A2'] = f"Адрес: {address}"
        ws['A4'], ws['B4'], ws['C4'], ws['D4'] = "Товар", "Цена", "Кол-во", "Итог"

        row = 5
        for item in cart_items:
            ws.cell(row=row, column=1, value=item.product.name)
            ws.cell(row=row, column=2, value=float(item.product.price))
            ws.cell(row=row, column=3, value=item.quantity)
            ws.cell(row=row, column=4, value=float(item.item_price))
            row += 1
        ws.cell(row=row+1, column=3, value="ИТОГО:")
        ws.cell(row=row+1, column=4, value=float(cart.total_price))

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        user_email = request.user.email if request.user.email else f"{request.user.username}@example.com"
        email = EmailMessage(
            subject=f"Ваш чек от Portative Shop (Заказ #{order.id})",
            body=f"Здравствуйте, {request.user.username}! Ваш чек находится во вложении.",
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[user_email],
        )
        email.attach(f"receipt_{order.id}.xlsx", excel_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        email.send()

        cart_items.delete()
        messages.success(request, f"Заказ #{order.id} успешно оформлен! Чек выслан на почту.")
        return redirect('profile')
    
    return render(request, 'shop/checkout.html', {'cart': cart})

def register_view(request):
    """Регистрация нового пользователя"""
    if request.method == 'POST':
        form = UserRegisterForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            messages.success(request, f'Добро пожаловать, {user.username}!')
            return redirect('home')
        else:
            for error in form.errors.values():
                messages.error(request, error)
    else:
        form = UserRegisterForm()
    
    return render(request, 'registration/register.html', {'form': form})

def login_view(request):
    """Вход в систему"""
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            messages.success(request, f'С возвращением, {user.username}!')
            next_url = request.GET.get('next', 'home')
            return redirect(next_url)
        else:
            messages.error(request, 'Неверное имя пользователя или пароль.')
    return render(request, 'registration/login.html')

def logout_view(request):
    """Выход из системы"""
    logout(request)
    messages.info(request, 'Вы вышли из системы.')
    return redirect('home')

@login_required
def profile_view(request):
    """Личный кабинет пользователя"""
    return render(request, 'shop/profile.html', {'profile': request.user.profile})

@login_required
def settings_view(request):
    """Страница настроек"""
    categories = Category.objects.all()
    return render(request, 'shop/settings.html', {'categories': categories})

# ==================== API VIEWSETS ====================

class CategoryViewSet(viewsets.ModelViewSet):
    queryset = Category.objects.all()
    serializer_class = CategorySerializer
    permission_classes = [IsAdminOrReadOnly]

class ManufacturerViewSet(viewsets.ModelViewSet):
    queryset = Manufacturer.objects.all()
    serializer_class = ManufacturerSerializer
    permission_classes = [IsAdminOrReadOnly]

class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.all()
    serializer_class = ProductSerializer
    permission_classes = [IsAdminOrReadOnly]
    
    def get_queryset(self):
        queryset = Product.objects.all()
        category = self.request.query_params.get('category', None)
        manufacturer = self.request.query_params.get('manufacturer', None)
        search = self.request.query_params.get('search', None)
        
        if category:
            queryset = queryset.filter(category_id=category)
        if manufacturer:
            queryset = queryset.filter(manufacturer_id=manufacturer)
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) | 
                Q(description__icontains=search)
            )
        return queryset

class CartViewSet(viewsets.ModelViewSet):
    serializer_class = CartSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        if self.request.user.is_staff:
            return Cart.objects.all()
        return Cart.objects.filter(user=self.request.user)
    
    def perform_create(self, serializer):
        serializer.save(user=self.request.user)

class CartItemViewSet(viewsets.ModelViewSet):
    serializer_class = CartItemSerializer
    permission_classes = [IsAuthenticated, IsOwnerOrAdmin]
    
    def get_queryset(self):
        return CartItem.objects.filter(cart__user=self.request.user)

class UserViewSet(viewsets.GenericViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [IsAuthenticated]
    
    @action(detail=False, methods=['get', 'patch'])
    def me(self, request):
        if request.method == 'GET':
            serializer = UserSerializer(request.user)
            return Response(serializer.data)
        
        elif request.method == 'PATCH':
            profile = request.user.profile
            serializer = ProfileSerializer(profile, data=request.data, partial=True)
            if serializer.is_valid():
                serializer.save()
                return Response(serializer.data)
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
    
    @action(detail=False, methods=['post'])
    def change_password(self, request):
        serializer = ChangePasswordSerializer(data=request.data)
        if serializer.is_valid():
            user = request.user
            if not user.check_password(serializer.data['old_password']):
                return Response({'error': 'Неверный старый пароль'}, status=status.HTTP_400_BAD_REQUEST)
            user.set_password(serializer.data['new_password'])
            user.save()
            update_session_auth_hash(request, user)
            return Response({'success': 'Пароль успешно изменен'})
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class OrderViewSet(viewsets.ModelViewSet):
    serializer_class = OrderSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        if self.request.user.is_staff:
            return Order.objects.all()
        return Order.objects.filter(user=self.request.user)
    
    def perform_create(self, serializer):
        serializer.save(user=self.request.user)
    
    @action(detail=True, methods=['patch'])
    def update_status(self, request, pk=None):
        if not request.user.is_staff:
            return Response({'error': 'Недостаточно прав'}, status=status.HTTP_403_FORBIDDEN)
        
        order = self.get_object()
        new_status = request.data.get('status')
        if new_status in dict(Order.STATUS_CHOICES):
            order.status = new_status
            order.save()
            return Response({'status': 'updated', 'new_status': new_status})
        return Response({'error': 'Неверный статус'}, status=status.HTTP_400_BAD_REQUEST)